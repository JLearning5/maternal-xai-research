from pathlib import Path
import sys

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
FILES = [
    ROOT / "chp~1~5" / "Adesanmi Anuoluwapo Gideon [chp~1~3~a].docx",
    ROOT / "chp~1~5" / "chp~4~5 [humanised~d].docx",
]

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if len(sys.argv) > 1:
    requested = sys.argv[1].lower()
    FILES = [path for path in FILES if requested in path.name.lower()]


for path in FILES:
    print(f"\n### {path.name}")
    document = Document(path)
    shown = 0
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        print(f"[{paragraph.style.name}] {text}")
        shown += 1
        if shown >= 180:
            break

    print(f"TABLES {len(document.tables)}")
    for table_index, table in enumerate(document.tables[:10]):
        print(f"TABLE {table_index}")
        for row in table.rows[:15]:
            print(" | ".join(" ".join(cell.text.split()) for cell in row.cells))
